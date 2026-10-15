from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from PIL import Image, ImageDraw, ImageFont
import json, math, zipfile

OUT = Path(__file__).resolve().parents[1] / '汇报材料'
OUT.mkdir(parents=True, exist_ok=True)
FONT = 'Noto Sans CJK SC'
C = dict(bg='F4F1E8', ink='183C36', dark='13352F', green='2D6956', light='E1E9DF',
         muted='62736D', sand='D8B87C', white='FFFFFF', orange='AD633F', pale='ECE3D3')

# These are established background/method references recalled from existing knowledge.
# Publisher/Crossref access was denied; no record is represented as newly read full text.
REFS = [
 dict(id='R1',author='Allen, C. D., et al.',year=2010,title='A global overview of drought and heat-induced tree mortality reveals emerging climate change risks for forests',journal='Forest Ecology and Management, 259, 660–684',doi='10.1016/j.foreco.2009.09.001',object='全球森林，多树种；非梭梭专门研究',question='干旱和高温相关树木死亡是否是广泛现象？',method='已报道死亡事件的全球性综述',result='综合讨论干旱、高温与森林死亡事件；提示开展跨区域、统一口径监测的必要性。',limit='事件汇总受报告偏差影响；不能推出河西走廊梭梭死亡率或阈值。',use='统一记录衰退、死亡时间与干旱背景；避免仅比较地区均值。'),
 dict(id='R2',author='McDowell, N., et al.',year=2008,title='Mechanisms of plant survival and mortality during drought: why do some plants survive while others succumb to drought?',journal='New Phytologist, 178, 719–739',doi='10.1111/j.1469-8137.2008.02436.x',object='植物干旱生理机制；跨物种',question='相同干旱下，为什么一些植物存活而另一些死亡？',method='机制综述与概念框架',result='以水力失效、碳限制及其相互作用解释干旱生存与死亡。',limit='机制框架不是梭梭现场验证；单项生理指标不足以诊断死亡原因。',use='提出水分供需失衡等竞争假说；低成本调查不能宣称证明碳饥饿。'),
 dict(id='R3',author='Anderegg, W. R. L., et al.',year=2016,title='Meta-analysis reveals that hydraulic traits explain cross-species patterns of drought-induced tree mortality across the globe',journal='Proceedings of the National Academy of Sciences, 113, 5024–5029',doi='10.1073/pnas.1525678113',object='全球多树种干旱死亡',question='哪些水力性状能解释物种间干旱死亡差异？',method='跨研究荟萃分析',result='水力性状对跨物种死亡差异具有解释价值。',limit='跨物种相关关系不等于梭梭种内可直接使用的风险阈值。',use='有仪器条件时增加水势等机制指标；必须在当地样本中验证。'),
 dict(id='R4',author='Adams, H. D., et al.',year=2017,title='A multi-species synthesis of physiological mechanisms in drought-induced tree mortality',journal='Nature Ecology & Evolution, 1, 1285–1291',doi='10.1038/s41559-017-0248-x',object='多树种干旱死亡生理研究',question='不同物种的死亡机制有哪些一致性和差异？',method='跨物种生理证据综合',result='将水力过程与碳相关过程置于共同框架，强调机制及相互作用。',limit='不能把木本植物的总体机制结果直接变成荒漠灌木管理标准。',use='将解释性变量和健康结局分开；机制指标采用嵌套精测。'),
 dict(id='R5',author='Cailleret, M., et al.',year=2017,title='A synthesis of radial growth patterns preceding tree mortality',journal='Global Change Biology, 23, 1675–1690',doi='10.1111/gcb.13535',object='死亡前径向生长轨迹；多树种',question='死亡前是否存在可识别的生长变化？',method='径向生长轨迹综合',result='死亡前生长轨迹具有信息，但信号因物种及背景而异。',limit='梭梭多干、偏心生长等条件下，年轮定年和代表性需专门验证。',use='借鉴重复生长监测；优先枝条增量和固定植株复测，不默认树钻可准确判龄。'),
 dict(id='R6',author='Choat, B., et al.',year=2018,title='Triggers of tree mortality under drought',journal='Nature, 558, 531–539',doi='10.1038/s41586-018-0240-x',object='树木干旱死亡机制综述',question='哪些过程触发干旱条件下的树木死亡？',method='生理机制综述',result='干旱死亡涉及水力、碳过程及生物因子的相互作用。',limit='不能据一次土壤含水率测定认定致死机制或通用阈值。',use='设置土壤水分、盐分、林分竞争和干扰等竞争解释。'),
 dict(id='R7',author='Breshears, D. D., et al.',year=2005,title='Regional vegetation die-off in response to global-change-type drought',journal='Proceedings of the National Academy of Sciences, 102, 15144–15148',doi='10.1073/pnas.0505734102',object='美国西南部干旱相关区域植被死亡，涉及松类',question='区域干旱冲击如何表现为大范围植被死亡？',method='地面观测与区域尺度资料结合',result='区域性死亡案例说明气候冲击需结合实地死亡观测解释。',limit='林型、物种及水文背景与梭梭不同；遥感变化也可能来自物候或背景。',use='用地面健康标签校验遥感，区分短期变黄与持续衰退。'),
 dict(id='R8',author='Roberts, D. R., et al.',year=2017,title='Cross-validation strategies for data with temporal, spatial, hierarchical, or phylogenetic structure',journal='Ecography, 40, 913–929',doi='10.1111/ecog.02881',object='具有空间、时间和层级结构的数据',question='相关样本应如何进行可信的模型验证？',method='交叉验证方法研究',result='验证划分应反映数据结构与模型的实际外推目标。',limit='留区验证也受区域数和事件数限制；不能只报告准确率。',use='按林分分组；若预测跨地区应用，补充留区域验证。'),
 dict(id='R9',author='Scheffer, M., et al.',year=2009,title='Early-warning signals for critical transitions',journal='Nature, 461, 53–59',doi='10.1038/nature08227',object='复杂系统临界转变理论',question='重大状态转变前是否存在早期信号？',method='理论及跨系统证据综述',result='提出临界减速等早期信号思路。',limit='需要合适的长时间序列；一次或两次普查不能验证临界减速。',use='作为远期监测启发；不纳入本硕士课题必做指标。'),
 dict(id='R10',author='Suding, K. N., Gross, K. L., & Houseman, G. R.',year=2004,title='Alternative states and positive feedbacks in restoration ecology',journal='Trends in Ecology & Evolution, 19, 46–53',doi='10.1016/j.tree.2003.10.005',object='恢复生态学；多种生态系统',question='退化状态为何可能持续，恢复为何并不自动发生？',method='理论与案例综合',result='反馈与替代状态框架提示：减弱胁迫不一定使生态系统自动恢复。',limit='不能仅凭存在枯枝就断言梭梭林已进入不可逆状态。',use='把自然更新与恢复潜力纳入分型；措施需小规模对照验证。'),
]

METRICS = [
 ['保存数量','现存活株密度','株/ha；活株数÷样方面积×10000','必做','用存活状态统一判定；死株、移除和补植另记','不把现存密度直接称作成活率'],
 ['保存数量','保存比例','现存同批活株÷该批初植株数×100%','有档案才做','初植数量、原始边界、补植与移除记录可核实','缺分母时仅报密度和档案缺失，不倒推比例'],
 ['个体健康','活冠/同化枝状态','固定角度照片；活冠比例%、枯梢比例%','必做','建立照片标尺；双人盲评部分照片','梭梭以同化枝为主，不能照搬阔叶树叶片黄化标准'],
 ['个体健康','当年生长量','标记枝条长度增量；cm/期','必做','固定枝位与季节；重复测量','不以株高大就等同于健康'],
 ['个体健康','生存和转级','活/死；健康等级；t0→t1','必做','永久标签、定位、照片；同季年度复测','只有同期分类时称现状诊断；未来风险需独立随访'],
 ['更新能力','实生苗与幼株','密度、大小级；标记苗存活比例','必做','区别萌蘖与实生苗；雨后及后续存活复查','单次出现幼苗不等于更新成功'],
 ['生态功能','林分冠盖与地表稳定性','冠盖度%；裸沙/固定地表比例%','必做代理项','统一样线/影像尺度；记录季节和地表类型','只能称功能代理，不能声称已测防风固沙量'],
 ['生态功能','风蚀/输沙或近地风速','集沙量、风速比等','精测选做','配对裸地/参照地；可比风况、传感器高度与时段','措施效果需前后对照；短期测值不外推年总量'],
 ['压力与解释','土壤含水率及盐分','质量或体积含水率；EC及测定方法','必做','分层、重复时点；同步记前期降雨；测容重时可换算储水','质量含水率与体积含水率不可直接混用'],
 ['压力与解释','林龄及初植密度','年；株/ha；造林档案可信等级','必做','档案/管理记录与访谈交叉核验；补植批次分开','林龄不直接作为退化分数；当前密度有幸存者偏差'],
 ['压力与解释','地下水和经营干扰','井水位或管理记录；放牧/采挖/灌溉史','可获则做','记录空间距离、时间与资料来源','无根系/同位素证据时不认定梭梭一定取用地下水'],
 ['机制精测','水势、气体交换或碳指标','依仪器规范和可比采样时点','选做','仅在匹配的健康/衰退子样本嵌套测定','不以单次测量证明水力失效或碳饥饿'],
]

SLIDES=[]
def add(section,title,sub='',cards=None,table=None,note='',source='研究设计建议；非实测结果',kind='normal'):
 SLIDES.append(dict(section=section,title=title,sub=sub,cards=cards or [],table=table,note=note,source=source,kind=kind))

add('开场','河西走廊人工梭梭林\n保存成效、衰退风险与分类管护','文献综述与硕士课题设计｜讨论初稿',
 cards=[('研究区域','武威 · 张掖 · 酒泉（含敦煌）'),('核心目标','看清现状 → 解释差异 → 识别风险 → 分类管护')],
 note='建议题目：河西走廊人工梭梭林保存成效及衰退风险的分层评价与分类管护。本稿按问题、思路、评价体系、方法、结果组织。重要边界：区域直接文献尚未完成数据库检索与全文核验；跨物种参考文献根据已有知识整理，书目信息及具体结论需回到原文复核。不能据此宣称河西走廊研究空白。',kind='cover')

add('问题','先把政府问题转成可检验的科学问题','同一套样地，贯通现状、解释、预警与管护。',cards=[
 ('01 保存得怎样？','现存数量是否维持？\n活着的植株是否健康？\n更新与防护功能是否持续？'),
 ('02 为什么不同？','在相近林龄和立地下，\n水分、盐分、密度及经营史\n能解释多少健康差异？'),
 ('03 哪些会继续衰退？','基线状态与暴露信息，\n能否识别随后出现新枯梢、\n生长下降或死亡的植株？'),
 ('04 应该怎么管？','按压力来源与恢复潜力分型，\n输出措施、适用条件、\n监测指标和优先级。')],note='“成效”需事先约定：如果评价防风固沙功能，就需直接测量或承认采用代理指标。当前保存比例不能等同于现存植株健康程度。硕士课题不必解释所有生理机制，但需要把观测、解释和预测的证据层次说清。')

add('问题','研究范围：三大片区，敦煌作为酒泉子区','地域不是唯一分层依据；造林历史和立地差异也必须进入设计。',cards=[
 ('武威片区','民勤等候选地点\n造林记录、地块边界和进入条件\n均须先向管理单位核实。'),
 ('张掖片区','临泽、高台等候选地点\n先清点可追溯的人工梭梭林，\n不预设每种林龄都存在。'),
 ('酒泉片区','肃州、金塔、瓜州、敦煌等候选地点\n敦煌可独立比较，\n统计汇总时避免与酒泉重复。')],note='这里列的是踏查候选区域，不是已核实样地清单。先建立总体地块名录，再依据立地、林龄及经营史选点。不能从少量便利样地推算整个城市的保存状况，更不能从几个地块直接外推河西走廊退化总面积。')

add('思路','这份综述的证据边界','先明确哪些内容已经能讨论，哪些尚不能下结论。',table=(['证据层','本版状态','可支持的用途'],[
 ['河西走廊人工梭梭林直接证据','尚未完成检索与全文核验','不能判断研究有无、区域覆盖和创新空白'],
 ['其他地区梭梭及近缘种研究','仅发现待核线索；未纳核心结论','用于后续追踪，不移植定量阈值'],
 ['其他木本植物衰退机制及方法','10篇背景/方法文献初步整理','构建竞争假说、监测与验证框架'],
 ['本课题评价与抽样方案','本报告提出；等待踏查与预调查','供导师讨论和方案收敛']]),note='本次 Crossref、OpenAlex、期刊及 DOI 网站读取均受网络策略限制，无法完成学术检索。已访问的 GitHub 数据仓库仅证明存在一个梭梭抗旱实验题名/数据线索，不是区域退化文献的替代。本稿不沿用用户已撤回的论文示例或其中数值。',source='检索状态说明；详见配套表“检索与核验”')

add('思路','文献应分四层读，避免混用证据','每篇都提取：问题—设计—指标—结果—限制—可借鉴点。',cards=[
 ('A 区域直接研究','河西走廊 × 人工梭梭林\n优先提取保存、衰退、林龄、\n水分与经营史。'),
 ('B 同物种研究','其他干旱区的人工/天然梭梭\n人工与天然分开；\n梭梭与白梭梭分开。'),
 ('C 其他植物退化','樟子松、柠条、油蒿、杨树等\n重点借鉴对照设计、\n时间监测及恢复试验。'),
 ('D 通用方法研究','健康评价、死亡机制、\n层级模型及分组验证\n用于提高论证可信度。')],note='层级A决定区域研究进展能否下结论，层级B用于物种相关假说，层级C用于方法借鉴，层级D用于分析规范。不能因为层级A没有检索完，就用全球森林文献来证明河西走廊研究不足。上述其他植物是建议扩展检索的对象，而不是本次已系统检索的结果。',source='综述组织方案；跨物种参考 R1–R10')

add('思路','其他植物研究：哪些成果最值得借鉴？','将跨物种结论转成可在梭梭上检验的设计。',table=(['代表文献','做到什么程度','本课题借鉴什么'],[
 ['R1 Allen 2010\nR7 Breshears 2005','汇总/观测区域性干旱死亡\n关注气候冲击与植被响应','地面死亡标签＋区域气候背景\n区分事件冲击与长期变化'],
 ['R2 McDowell 2008\nR6 Choat 2018','形成水力、碳过程及相互作用框架','并列检验水分、盐分和竞争解释\n不凭一个指标判机制'],
 ['R3 Anderegg 2016\nR4 Adams 2017','比较跨物种生理差异与死亡机制','部分样地嵌套精测\n不套用别的物种致死阈值'],
 ['R5 Cailleret 2017','综合死亡前生长轨迹','固定植株的重复测量\n检验基线指标能否预测后续变化']]),note='这些是成熟的机制及研究设计思路，并不构成梭梭区域结果。全部参考资料在本次未能访问出版者或全文，概括保持在基础主题层面，不报告未核对的效应量、样本量或死亡阈值。',source='R1–R7；既有文献主题整理，本次未复核全文')

add('思路','评价“做到什么程度”，看证据而非模型数量','后续读区域文献时，逐项判断它真正跨过了哪一步。',table=(['层次','能回答的问题','还不能据此声称'],[
 ['Ⅰ 单次样地描述','现在保存和健康状况如何？','未来会不会死亡'],
 ['Ⅱ 匹配对照与分层分析','哪些因素与现状差异相关？','已证明因果及唯一机制'],
 ['Ⅲ 同株随访＋分组验证','哪些基线信息能识别后续变化？','能准确预测任意地区与年份'],
 ['Ⅳ 管护试验＋前后对照','某措施在何种条件下有效？','已适用于整个河西走廊']]),note='这是本报告提出的阅读评价尺，不是对尚未读取的区域文献作出的分级结论。评价时同时记录林分独立重复、样本选择、混杂控制、验证集独立性和不确定性。复杂模型不能补救无重复或无时间标签的设计。')

add('思路','你的潜在推进点：把四个环节接起来','优先做一个可信闭环，再考虑扩展机制和遥感。',cards=[
 ('统一口径','相同的保存与健康定义，\n支持跨片区、跨林龄比较。'),
 ('提高可比性','同龄不同地＋同地不同龄，\n有独立林分重复；\n明确无法拆开的混杂。'),
 ('加入时间','同株、同季随访，\n用新发生衰退验证基线指标。'),
 ('落实措施','状态 × 压力 × 恢复潜力，\n配套分型、监测和试验。')],note='这些应表述为“拟开展的改进”，是否超越已有河西走廊梭梭研究，必须完成直接文献的系统检索后再定。若已有研究完成前三项，可进一步聚焦外区验证、低成本指标的独立预测价值，或一个有对照的管护试验。不得承诺“首次”。')

add('评价体系','先分清：保存、健康、功能与风险','同一林分可能保存比例较高，但活冠下降、更新失败。',cards=[
 ('保存','原来种下的植株，\n如今还剩多少？\n核心：同批活株＋可靠分母。'),
 ('健康','现存植株是否持续生长？\n核心：活冠、枯梢、\n当年生长及状态变化。'),
 ('功能','是否维持预定防护作用？\n核心：覆盖和地表代理，\n有条件再测输沙或风速。'),
 ('风险','未来是否会继续变差？\n核心：后续新发结局，\n用基线信息独立验证。')],note='避免把这四项混为一个随意加权总分。政府需要看见具体短板：数量下降、活力下降、更新失败或功能不足，可能对应不同管护措施。风险必须有明确的时间窗口与结局定义。')

add('评价体系','核心指标：保存数量与个体健康','优先选择可重复、可审查、成本可承受的指标。',table=(['维度','建议核心指标','关键口径'],[
 ['保存数量','活株密度；同批保存比例（有档案时）','补植和移除分开；原始边界一致\n无初植分母时不报“成活率”'],
 ['冠层健康','活冠比例；枯梢比例；标准照片','依据梭梭同化枝特点建立评分标尺\n抽取照片双人盲评'],
 ['生长活力','固定枝条长度增量；必要时基径增量','同枝位、同季节、同测量方式\n多干结构单独记录'],
 ['状态转变','同株健康等级转移；新增死亡','永久标签＋定位＋影像\n先定义失访和机械损伤的处理']]),note='保存比例=现存同批活株数/初植株数。若发生补植，应按造林批次重建分母。基线已有死亡不能用于验证未来死亡预测。个体健康标签应主要来自症状和生长，水分与盐分作为解释变量保留在标签外，避免循环论证。')

add('评价体系','更新和功能必须有位置，压力指标另列','测了土壤含水率，不等于测了防风固沙成效。',table=(['角色','指标建议','解释边界'],[
 ['更新能力','实生苗/幼株密度与大小级\n标记苗的后续存活','萌蘖与实生苗分开\n一次萌发≠成功更新'],
 ['功能代理','林分冠盖度、裸沙比例\n地表固定状况','可用于功能相关状态评价\n不能直接换算固沙量'],
 ['功能精测（选做）','配对地表的风速/输沙观测','需控制风况、地形、测高和时段\n不可凭单日数据算年效益'],
 ['压力与背景','分层土壤水分、盐分、林龄\n初植密度、经营史及干扰','解释差异或形成假说\n不直接纳入健康结局定义']]),note='生态功能是否列为主线，应根据仪器和人力确定。没有直接功能测量，论文题目和结论要使用“功能相关指标”或“防护状态代理”。地下水资料可作为环境背景，不能无证据地推断梭梭直接依赖地下水。',source='研究设计建议；机制框架参考 R2、R6、R10')

add('评价体系','分级规则：先有症状标尺，再谈综合评分','当前没有经过本地验证的统一梭梭退化阈值可直接采用。',cards=[
 ('第一步｜预调查定标','建立同化枝和冠层照片库；\n定义健康、轻度、明显衰退、\n死亡等可观察状态。'),
 ('第二步｜一致性检查','双人独立评分；\n检查等级一致性，\n对争议植株现场复核。'),
 ('第三步｜外部检验','检验分级是否对应随后生长\n或健康恶化；若有总分，\n比较不同权重的敏感性。')],note='这里不预设20%、50%等等级分界。阈值需参考经过核实的规范/文献，并在预调查中检验可操作性。死亡需要明确无存活组织的判定和必要复查。若使用PCA或熵权，必须解释权重含义；统计权重不自动具有生态意义。')

add('方法','抽样：总体调查与机制精测分开','估计现状比例需要代表性；解释机制需要可比性。',cards=[
 ('总体样本｜描述现状','先建人工林地块名录；\n按区域、立地和林龄分层抽取，\n记录抽样概率及无法进入的地块。'),
 ('匹配子样本｜解释差异','在相近林龄、土壤与经营史下，\n比较健康与衰退林分；\n补充水分和生理精测。'),
 ('同株随访｜识别风险','固定植株基线调查，\n同季节复测健康和生长；\n有年度复测更适合判断持续衰退。')],note='若有意增加明显衰退样地，不能直接把样本中退化比例当作区域总体比例，必须使用设计权重或分别报告。便利抽样只能描述样本地块。季内复测可以看生长及短期变化，但不能自动证明持续退化或跨年死亡风险。')

add('方法','样本量示例：27个林分，54个样方','这是工作量示例，最终以可用林分、预调查变异和失访率调整。',table=(['层级','示例配置','统计含义'],[
 ['区域 × 林龄层 × 独立林分','3片区 × 3林龄层 × 3林分 = 27林分','林龄层据档案和分布确定\n争取各片区均有可比年龄'],
 ['林分内调查','每林分2个样方，共54样方','按低密林预调查调整面积\n样方不替代独立林分'],
 ['固定植株复测','每样方标记约10株，约540株','需覆盖状态和大小差异\n实际事件数决定模型复杂度'],
 ['精测子集','各片区约3个林分，共约9林分','匹配健康/衰退状态\n统一采样季节与时间窗']]),note='例示样本量不是功效计算结果。只有27个林分，不宜同时拟合大量林分层变量或复杂机器学习模型。若死亡/恶化事件很少，应简化目标或延长随访，优先报告状态变化与置信区间。不能把540株当成540个相互独立的林分。')

add('方法','解释差异：四组竞争假说','把可测因素放在同一个框架中，避免“相关即原因”。',cards=[
 ('H1 水分供需失衡','分层水分与前期降雨\n是否解释同立地的生长下降？\n有条件测水势验证。'),
 ('H2 密度与竞争','控制林龄、立地和经营史后，\n初植密度/邻域竞争是否\n与后续恶化相关？'),
 ('H3 立地与盐分','土壤质地、盐分和地貌\n是否改变水分与健康之间\n的关系？'),
 ('H4 经营与干扰','补水、封育、放牧、采挖、\n虫害及补植史，\n能否解释剩余差异？')],note='当前密度可能因先前死亡而降低，直接把它用于解释当前退化会出现反向因果。初植密度和邻域竞争也需明确时间来源。林龄若和地区/生境完全重合，统计模型无法凭空拆开影响；应改进抽样或承认证据限制。',source='本课题待检验假说；参考 R2、R4、R6')

add('方法','分析路线：从可信描述开始','每一步都对应一个研究问题，不为使用模型而加模型。',table=(['目标','优先方法','必须检查'],[
 ['保存与健康现状','分层描述；林分级汇总\n置信区间；设计加权（适用时）','缺档案率、失访率、样本覆盖\n空间与层级相关性'],
 ['差异解释','少量预设变量的混合模型\n林分/样方层级及匹配比较','共线性、混杂、响应分布\n关联解释与因果声明分开'],
 ['后续变化识别','随访结局的简约回归\n有事件时间再考虑生存模型','事件数、删失/失访\n时间先后和预测变量泄漏'],
 ['空间扩展（选做）','地面标签＋遥感指标\n按林分/区域分组验证','裸土和物候干扰\n适用范围及预测不确定性']]),note='活冠比例应采用适合比例数据的模型或变换；等级结局可用有序模型；不要把所有结局都机械套正态回归。变量数由独立林分数和事件数约束。相关性强且同处因果链的指标不宜不加判断同时入模。',source='分析方案；验证策略参考 R8')

add('方法','风险识别：先定义“未来发生了什么”','若只有一次调查，只能给出当前脆弱性，不能给出已验证未来风险。',cards=[
 ('结局标签','在明确随访期内：\n健康等级恶化、持续生长下降\n或新发生死亡；分别报告。'),
 ('预测信息','只使用基线可获取的信息；\n症状、水分、盐分和经营背景\n不使用随访期结果作特征。'),
 ('验证标准','按林分分组划分训练与验证；\n报告校准、召回率和误报，\n事件少时只做探索分析。')],note='预测当期健康等级并非预测未来衰退。若响应变量由枯梢比例构造，再把同一次枯梢比例作为特征，可得到高准确率但没有预警价值。验证需要所有预处理、变量筛选和阈值选择仅在训练折中进行；同一植株及同林分样本不能随机分散导致泄漏。',source='研究设计建议；参考 R5、R8')

add('方法','复杂方法的取舍：以硕士工作量为边界','核心任务做到可重复；其余根据问题和条件增加。',table=(['层级','建议安排','理由'],[
 ['必做','档案核查、独立林分重复\n统一健康调查、同株复测\n水分/盐分与经营史','直接支撑现状、差异和后续变化\n成果可被复核'],
 ['择一加强','水势/气体交换精测\n或一个小规模管护试验\n或有标签的遥感扩展','服务一条清楚的科学假说\n取决于仪器、时间和事件数'],
 ['暂不必做','同时上同位素、宏基因组\n复杂结构方程和多套机器学习\n无长时序的临界转变检测','容易分散工作量\n不能替代对照、重复及时间信息']]),note='“更好”优先体现在证据质量，而不是指标和模型的数量。若不能实现跨年复测，就应将论文目标调整为“现状诊断与脆弱性分型”，避免题目写成未来衰退预测。R9是长时序早期信号思路，不应在一次普查中强行使用。',source='工作量建议；R9仅作远期方法背景')

add('结果','文献主题归纳：哪些认识可以作为出发点？','以下为跨物种背景认识；尚未在本课题梭梭样地中得到验证。',table=(['背景认识','参考','能支持什么 / 不能支持什么'],[
 ['干旱死亡需要区域与时间背景','R1、R7','支持记录气候与健康变化\n不支持本地死亡比例估计'],
 ['水力、碳过程及生物因素可能交互','R2、R4、R6','支持多假说与嵌套精测\n不支持单指标判定机制'],
 ['死亡前生长轨迹可能提供信息','R5','支持同株重复监测\n不支持直接套用梭梭预警阈值'],
 ['验证必须符合空间与层级结构','R8','支持按林分/区域分组验证\n不保证小样本模型可稳定外推']]),note='这是文献主题层面的归纳，不是本次系统综述的定量合并结果。未统计区域论文篇数、没有做效应量合并，也未复核每篇全文。不能用“共识”替代针对梭梭的当地验证。',source='R1、R2、R4–R8；既有知识整理，待全文复核')

add('结果','课题最终应交付四类结果','这是预期成果结构；目前没有本地实测数据或结果图。',cards=[
 ('现状清单与图表','各分层的活株密度、健康分布、\n更新及功能代理；\n同时标明抽样覆盖和不确定性。'),
 ('差异解释','主要关联因素及效应区间；\n不同立地/林龄的差异，\n以及无法拆开的混杂。'),
 ('风险或脆弱性分型','有随访：验证未来恶化识别；\n无随访：仅报告脆弱性，\n说明适用范围。'),
 ('政府管护清单','具体地块、问题类型、\n措施条件、优先级、\n监测要求与试验建议。')],note='不绘制虚构退化率、风险热力图或处理效果。后续图表应包括样本流程图、林分健康分布图、匹配比较或效应图、真实随访转移图和管护分区图。区域总体估计必须有足够代表性和合法地块边界支持。')

add('结果','分类管护：从问题类型到有条件的措施','以下是管理决策框架，措施成效仍需当地试验和复测。',table=(['类型','识别依据','优先行动','验证要求'],[
 ['相对稳定','健康维持且更新可持续','保护＋常规巡查','同季复测，关注干旱年'],
 ['健康但暴露较高','当前活力尚可；水分/盐分压力升高','加密监测；减少新增干扰','不能等同已验证高死亡概率'],
 ['竞争相关衰退','可比样地支持竞争解释','小规模疏密对照试验','同步看水分、健康和裸沙风险'],
 ['立地/干扰相关衰退','盐分、风蚀或经营记录有证据','治理对应压力；有条件封育','前后对照；兼顾地表稳定'],
 ['更新受限或重度衰退','幼苗难维持、活力低且功能下降','先查水资源与立地\n再决定补植或结构调整','避免盲目高密补植；分阶段试验']]),note='不建议仅凭土壤含水率低就统一灌溉，或仅凭林龄大就统一砍伐更新。疏密、补植和灌溉均涉及管理审批、生态权衡与成本，不是本稿已验证的技术处方。风险分级是监测与优先级工具，不能替代经营决策。',source='研究设计与管理建议；恢复思路参考 R10')

add('结果','如果要证明措施有效，至少保留一个对照','一个小而规范的试验，比没有对照的大规模建议更有说服力。',cards=[
 ('设计','选择证据支持的一种措施；\n设处理与对照，尽量随机分配；\n实施前先测基线。'),
 ('观测','同步复测健康、生长与水分；\n若疏密或地表操作，\n加测裸沙/风蚀风险。'),
 ('分析','比较处理与对照的变化差异；\n统计单位对应处理分配单位；\n报告区间、成本和副作用。')],note='处理分配给整个林分时，林分内几十株植物仍不是几十个独立处理重复。少数试验地块的结果应作为试点证据，不能直接推广全走廊。若论文时间不足，可将试验列为下一阶段验证，并将当期建议明确限定为待验证方案。')

add('思路','文献检索：接下来需要补齐的直接证据','不只检索“退化”；“保存率、生长衰退、死亡、土壤干化”也要覆盖。',table=(['检索层','中文关键词组','英文关键词组'],[
 ['区域与物种','河西走廊/民勤/武威/张掖/临泽/\n高台/酒泉/金塔/瓜州/敦煌＋梭梭','Haloxylon ammodendron +\nHexi / Minqin / Zhangye /\nJiuquan / Dunhuang'],
 ['人工林与现状','人工林/造林/保存率/存活/\n死亡/枯梢/衰退/退化','plantation / afforestation /\nsurvival / mortality / decline'],
 ['解释与管理','林龄/密度/土壤水分/干化/\n盐分/地下水/抚育/疏伐/更新','stand age / density / soil water /\nsalinity / thinning / regeneration'],
 ['其他植物借鉴','樟子松/柠条/油蒿/杨树＋\n衰退/死亡/预警/复壮','dieback / drought mortality /\nvitality / restoration']]),note='建议数据库：中国知网、万方、Web of Science/Scopus（学校授权）、Crossref和出版者全文。按数据库语法调整检索式；不虚报检索日期和命中篇数。正文中的植物拉丁名、研究区域和人工/天然林属性需要人工复核。',source='可执行检索方案；配套表中提供可复制检索式')

add('思路','每读一篇，就回答这九个问题','这张表决定能否判断“别人做到哪一步、我们补什么”。',cards=[
 ('研究设计','① 在哪里、什么林？\n② 林龄和造林记录可靠吗？\n③ 有多少独立林分和时间点？'),
 ('评价与结果','④ 退化如何定义和测量？\n⑤ 实际结果是什么、误差多大？\n⑥ 有无时间顺序及混杂控制？'),
 ('借鉴与推进','⑦ 有没有独立验证？\n⑧ 措施有效性是否有对照？\n⑨ 本课题具体补哪一环？')],note='配套Excel中把“未报告”“未核实”和“没有”分开填写。全文没有提供样本量不能记作0，检索未命中不能写作不存在。若只读摘要，不填写只有全文才能确认的取样重复和机制证据。')

add('结论','硕士论文的主线：评价—解释—验证—建议','建议先收敛到一条能按时完成、可被复核的证据链。',cards=[
 ('论文题目建议','河西走廊人工梭梭林保存成效\n及衰退风险的分层评价与分类管护\n（有随访时使用“风险”）'),
 ('资源不足时','人工梭梭林健康状况、\n关联因素及脆弱性分型\n（单次调查的稳妥定位）'),
 ('真正要守住的底线','代表性＋独立林分重复\n指标与结局分开＋真实时间标签\n结论范围与证据匹配')],note='建议论文正文按“区域与设计—现状评价—影响因素与后续变化—分类管护”组织。政府建议落到地块清单与监测表；学术贡献落到可靠的对照、时间证据和验证。当前仍需补齐河西走廊直接文献，才可以最终确定创新点与论文题名。')

add('参考文献','背景与方法文献 01—05','书目信息依据已有知识整理；本次网络限制，未复核出版者记录与全文。',table=(['编号 / 作者','题名','期刊与 DOI'],[
 [r['id']+' / '+r['author']+' '+str(r['year']),r['title'],r['journal']+'\n'+r['doi']] for r in REFS[:5]]),note='全部文献主题及适用边界见配套Excel。可以通过DOI定位原文，但本次无法访问DOI或出版者进行核验。不得把这些通用文献写作河西走廊人工梭梭林的直接调查证据。',source='R1–R5；配套Excel含可点击 DOI 链接')

add('参考文献','背景与方法文献 06—10','它们用于构建方法框架，不能替代河西走廊的物种与区域证据。',table=(['编号 / 作者','题名','期刊与 DOI'],[
 [r['id']+' / '+r['author']+' '+str(r['year']),r['title'],r['journal']+'\n'+r['doi']] for r in REFS[5:]]),note='R9的临界转变方法要求合适的时间序列；本课题不以两次调查宣称检测临界减速。R10用于恢复潜力思考，不能据它认定当地林分不可逆退化。',source='R6–R10；配套Excel含可点击 DOI 链接')

add('附录','本版已交付什么，还差什么？','用于选题与导师讨论的框架稿；区域专门文献综述尚未完成。',cards=[
 ('已交付','可编辑PPT＋演讲备注\n背景/方法文献对照表\n指标表、抽样示例与检索式'),
 ('仍需补齐','河西走廊直接文献题录与全文\n逐篇核验样本、方法和真实结果\n据此修订创新点与阈值'),
 ('下一次收敛','结合造林档案和踏查可行性\n确定代表区域、随访周期\n以及一个最值得加强的方法')],note='本版不包含虚构调查结果，不声称完成系统综述。用户之前撤回的论文示例及数值未用作核心依据。后续获得PDF/题录或恢复数据库访问后，应优先替换区域直接证据部分，再复核背景参考资料。')

prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
def rgb(h):return RGBColor.from_string(C.get(h,h))
def rect(sl,x,y,w,h,fill,line=None,radius=False):
 sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
 sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
 if line:sh.line.color.rgb=rgb(line)
 else:sh.line.fill.background()
 if radius:
  try: sh.adjustments[0]=.06
  except:pass
 return sh
def text(sl,x,y,w,h,content,size=20,color='ink',bold=False):
 # Fit to the actual text rectangle using CJK font metrics before saving.
 # PowerPoint text overflow is otherwise silent, particularly with mixed CJK/Latin.
 fpath='/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc' if bold else '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
 while size>8:
  ft=ImageFont.truetype(fpath,round(size*2))
  maxwidth=w*144
  lines=0
  for line in content.split('\n'):
   buf='';cnt=1
   for char in line:
    if ft.getlength(buf+char)>maxwidth and buf:
     cnt+=1;buf=char
    else:buf+=char
   lines+=cnt
  required=lines*size*1.16+max(0,len(content.split('\n'))-1)*3
  if required<=h*72-2:break
  size-=.5
 sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame
 tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
 for j,line in enumerate(content.split('\n')):
  p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=line
  p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
  p.space_after=Pt(3);p.line_spacing=1.10
 return sh
def base(i,s):
 sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=rgb('bg')
 rect(sl,.42,.39,.08,.3,'green');text(sl,.65,.38,10,.3,s['section']+'  /  HALOXYLON REVIEW',10,'green',True)
 text(sl,.65,.94,12.03,.7,s['title'],27,'ink',True)
 text(sl,.65,1.77,12,.56,s['sub'],14,'muted')
 rect(sl,.65,6.94,12.03,.012,'sand')
 text(sl,.65,7.08,11.4,.22,s['source'],8,'muted')
 text(sl,12.05,7.04,.65,.26,f'{i:02d} / {len(SLIDES):02d}',9,'green',True)
 sl.notes_slide.notes_text_frame.text=s['note']+'\n\n来源/状态：'+s['source']
 return sl

for i,s in enumerate(SLIDES,1):
 if s['kind']=='cover':
  sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=rgb('dark')
  # Quiet vector landscape, deliberately schematic and not a geographic map.
  for k in range(6):
   sh=sl.shapes.add_shape(MSO_SHAPE.OVAL,Inches(6.7+k*.25),Inches(4.8+k*.26),Inches(9),Inches(4))
   sh.fill.solid();sh.fill.fore_color.rgb=rgb(['green','183E34','285044','4B6652','79866A','D8B87C'][k]);sh.line.fill.background()
  text(sl,.85,.7,10,.4,'HEXI CORRIDOR  /  RESEARCH REVIEW',12,'sand',True)
  text(sl,.85,1.65,11.5,1.7,s['title'],33,'white',True)
  text(sl,.9,3.65,11,.5,s['sub'],18,'sand')
  text(sl,.9,4.6,10,.5,s['cards'][0][1],20,'white')
  text(sl,.9,5.38,11,.5,s['cards'][1][1],19,'white')
  text(sl,.9,6.8,11,.3,'范围说明：区域直接文献待核验；结果页区分背景认识与预期成果。',11,'sand')
  sl.notes_slide.notes_text_frame.text=s['note'];continue
 sl=base(i,s)
 if s['table']:
  hdr,rows=s['table'];cols=len(hdr);n=len(rows)
  widths=([2.3,4.2,5.53] if cols==3 else [2.0,3.05,3.5,3.48])
  if s['section']=='参考文献':widths=[2.05,6.15,3.83]
  th=.51;rowh=min(.88,(4.2-th)/n);top=2.45
  fs=14 if n<=4 else 12
  if s['section']=='参考文献':fs=10.2;rowh=.73
  xx=.65
  for j,h in enumerate(hdr):
   rect(sl,xx,top,widths[j]-.015,th,'green');text(sl,xx+.12,top+.12,widths[j]-.24,.3,h,13,'white',True);xx+=widths[j]
  for k,row in enumerate(rows):
   xx=.65;y=top+th+k*rowh
   for j,cell in enumerate(row):
    rect(sl,xx,y,widths[j]-.015,rowh-.018,'white' if k%2==0 else 'light')
    text(sl,xx+.12,y+.12,widths[j]-.24,rowh-.19,cell,fs,'ink',j==0);xx+=widths[j]
 else:
  cards=s['cards'];n=len(cards);gap=.2
  if n==4:
   w=5.9;h=1.88
   coords=[(.65,2.48),(6.76,2.48),(.65,4.55),(6.76,4.55)]
   for k,((head,body),(x,y)) in enumerate(zip(cards,coords)):
    rect(sl,x,y,w,h,'white',radius=True);rect(sl,x,y,.055,h,'sand')
    text(sl,x+.23,y+.2,w-.46,.35,head,19,'green',True)
    text(sl,x+.23,y+.7,w-.46,1.10,body,16.5,'ink')
  else:
   w=(12.03-gap*(n-1))/n;h=3.85
   for k,(head,body) in enumerate(cards):
    x=.65+k*(w+gap);y=2.5
    rect(sl,x,y,w,h,'white',radius=True);rect(sl,x+.22,y+.3,.44,.055,'sand')
    text(sl,x+.23,y+.64,w-.46,.7,head,19,'green',True)
    text(sl,x+.23,y+1.56,w-.46,2.05,body,18,'ink')

prs.core_properties.title='河西走廊人工梭梭林保存成效、衰退风险与分类管护'
prs.core_properties.subject='文献主题梳理与硕士课题设计；区域证据待核验'
prs.core_properties.author=''
pptx_path=OUT/'河西走廊人工梭梭林_文献梳理与课题设计.pptx';prs.save(pptx_path)

# Workbook: actual references, honest screening status, and field-ready design.
wb=Workbook();ws=wb.active;ws.title='阅读说明'
for row in [
 ['项目','说明'],
 ['用途','硕士选题、导师讨论及文献阅读工作底稿；不是已完成的区域系统综述'],
 ['核心范围','武威、张掖、酒泉人工梭梭林；敦煌属于酒泉，可作为子区'],
 ['本次检索限制','Crossref/OpenAlex/DOI和期刊请求被网络策略拒绝，未完成直接文献检索与全文核验'],
 ['参考文献状态','R1–R10来自已有知识；本次未联网复核，不填写未核实效应量或样本量'],
 ['区域证据状态','没有已核验的区域直接论文纳入本版；不等于该区域不存在相关研究'],
 ['撤回内容','用户先前发错的论文示例及其中数值均未纳入核心证据'],
 ['结果含义','PPT中区分跨物种背景认识、待检验假说与预期成果；无虚构调查结果'],
 ['样本量状态','27林分、54样方、约540标记株只是工作量示例；需预调查和功效/事件数评估'],
 ['下一步','获取直接文献全文与造林档案，完成证据矩阵，再确认创新点与现场方案']]:ws.append(row)

ws=wb.create_sheet('背景与方法文献')
ws.append(['编号','作者','年份','题名','期刊','DOI','DOI链接','对象与区域','问题','方法','结果主题（非全文摘录）','局限与外推边界','本课题借鉴','核验状态','证据层'])
for r in REFS:
 ws.append([r['id'],r['author'],r['year'],r['title'],r['journal'],r['doi'],'https://doi.org/'+r['doi'],r['object'],r['question'],r['method'],r['result'],r['limit'],r['use'],'依据已有知识整理；本次未核验出版者记录/全文','跨物种背景或通用方法，非河西直接证据'])
 ws.cell(ws.max_row,7).hyperlink='https://doi.org/'+r['doi'];ws.cell(ws.max_row,7).style='Hyperlink'

ws=wb.create_sheet('区域文献待补')
ws.append(['待检索主题','建议关键词','必须抽取的字段','纳入条件','当前状态'])
for row in [
 ['保存成效','河西走廊/民勤/武威/张掖/酒泉/敦煌＋梭梭＋人工林/保存率','初植分母、补植/移除、现存密度、边界、造林记录','原始研究；能区分人工梭梭和其他林型','尚未完成数据库检索与全文核验'],
 ['衰退及水分','梭梭＋衰退/枯梢/死亡/土壤干化/盐分','林龄×立地组合、独立重复、水分深度和时间、健康结局','提供可辨认的衰退/生长/死亡或水分证据','待检索'],
 ['风险与动态','梭梭＋动态/复测/预警/风险/死亡预测','随访期、个体标签、失访、独立验证、事件数','明确时间顺序；区分当前诊断与未来预测','待检索'],
 ['管护恢复','梭梭＋疏伐/密度调控/封育/补植/更新/复壮','处理分配单位、对照、基线、复测、成本、副作用','区分经验建议和对照试验结果','待检索']]:ws.append(row)

ws=wb.create_sheet('评价指标与口径');ws.append(['角色/维度','指标','单位/定义','优先级','实施要求','避免的误读'])
for row in METRICS:ws.append(row)

ws=wb.create_sheet('检索与核验');ws.append(['类别','数据库/来源','检索式或链接','本次状态','可支持的结论'])
queries=[
 ('区域中文','CNKI/万方','(梭梭 OR 人工梭梭林) AND (河西走廊 OR 民勤 OR 武威 OR 张掖 OR 临泽 OR 高台 OR 酒泉 OR 金塔 OR 瓜州 OR 敦煌)'),
 ('区域主题','CNKI/万方','梭梭 AND (保存率 OR 成活率 OR 生长 OR 枯梢 OR 衰退 OR 退化 OR 死亡 OR 土壤干化 OR 更新)'),
 ('区域英文','Web of Science/Scopus','("Haloxylon ammodendron") AND ("Hexi Corridor" OR Minqin OR Wuwei OR Zhangye OR Linze OR Gaotai OR Jiuquan OR Jinta OR Guazhou OR Dunhuang)'),
 ('物种英文','Web of Science/Scopus','("Haloxylon ammodendron") AND (plantation* OR afforestation) AND (decline OR dieback OR mortality OR survival OR "soil water" OR regeneration)'),
 ('近缘种扩展','Web of Science/Scopus','("Haloxylon persicum") AND (plantation* OR mortality OR drought OR regeneration)'),
 ('其他植物中文','CNKI/万方','(樟子松 OR 柠条 OR 油蒿 OR 杨树) AND (人工林 OR 植被) AND (衰退 OR 退化 OR 死亡 OR 风险 OR 复壮)'),
 ('方法英文','Web of Science/Scopus','("tree mortality" OR "shrub dieback") AND (longitudinal OR "repeated measurement" OR "spatial cross-validation" OR restoration)'),
]
for a,b,c in queries:ws.append([a,b,c,'计划检索；非已执行命中结果','用于后续构建区域直接证据；语法按数据库调整'])
for a,b in [('Crossref','https://api.crossref.org/works'),('OpenAlex','https://api.openalex.org/works'),('DOI','https://doi.org/'),('中国沙漠','https://www.desert.ac.cn/'),('生态学报','https://www.ecologica.cn/')]:
 ws.append(['实际连接检查',a,b,'本次请求被网络代理拒绝（403）','不能获取题录、摘要或全文；不能判断文献不存在'])
ws.append(['数据线索','GitHub公开仓库','https://github.com/5100163/Original-Data','已读取仓库页；仅确认仓库题名和文件名，未核对论文','Dark septate endophytes enhance the drought tolerance of Haloxylon ammodendron in sterilized and nonsterilized soil；仅作抗旱实验线索，不作为区域退化结论'])

ws=wb.create_sheet('现场与分析方案');ws.append(['环节','建议','质量控制','证据边界'])
for row in [
 ['总体名录','地块ID、边界、人工/天然、造林批次、初植密度、经营史','档案/访谈交叉核验；标记信息可信等级','名录不全时不作完整区域普查推断'],
 ['分层抽样','区域×立地×林龄；随机或概率抽样','避免选择全是明显衰退样地；记录无法进入地块','便利样地比例≠区域退化比例'],
 ['样本示例','27独立林分、每林分2样方、每样方约10标记株','依预调查变异、实际林分分布及事件数调整','不是既定方案或功效分析结果'],
 ['健康标签','活冠/枯梢、固定枝生长、健康转级和存活','照片标尺、双人盲评、复查、明确失访','基线症状不直接当作未来结局'],
 ['水分盐分','按土层和时点重复测定；记录降雨、质地及EC方法','含水率单位统一；必要时测容重','瞬时值不等同长期胁迫；不预设2%通用阈值'],
 ['解释模型','少量预设变量；控制层级、林龄、立地和经营史','检查共线性和反向因果；报告效应区间','观察关联不自动等于因果机制'],
 ['预测验证','基线信息预测随后新发结局；按林分分组验证','全部筛选/预处理在训练折内；报告校准与误报','事件少时简化；一次调查仅脆弱性'],
 ['措施试验','选择一类措施；处理/对照；测前测后','重复与随机化对应处理分配单位','只有建议无试验不能称措施已有效'],
 ['政府输出','地块问题清单、分类措施、优先级及监测要求','给出证据等级、不确定性与适用条件','不直接作统一疏伐/灌溉/高密补植处方']]:ws.append(row)

for ws in wb:
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for cell in ws[1]:cell.fill=PatternFill('solid',fgColor=C['green']);cell.font=Font(name='Noto Sans CJK SC',color='FFFFFF',bold=True,size=11)
 for row in ws.iter_rows(min_row=2):
  for cell in row:
   cell.alignment=Alignment(vertical='top',wrap_text=True);cell.font=Font(name='Noto Sans CJK SC',size=10,color=C['ink'])
   if cell.row%2==0:cell.fill=PatternFill('solid',fgColor='EAF0E7')
 for col in ws.columns:
  letter=col[0].column_letter;ws.column_dimensions[letter].width=36 if ws.title!='背景与方法文献' else 34
 ws.column_dimensions['A'].width=23
 for r in range(2,ws.max_row+1):ws.row_dimensions[r].height=86 if ws.title=='背景与方法文献' else 70
 ws.row_dimensions[1].height=30
ws=wb['背景与方法文献'];ws.column_dimensions['D'].width=70;ws.column_dimensions['G'].width=48
xlsx_path=OUT/'梭梭林文献矩阵_评价指标与检索计划.xlsx';wb.save(xlsx_path)

md=['# 河西走廊人工梭梭林：文献梳理与硕士课题设计\n',
 '> 讨论初稿。区域直接文献尚未完成检索与全文核验；以下不能用来宣称研究空白、首次发现或区域退化率。\n',
 '本稿使用用户更正后的课题方向，不使用先前发错的论文示例与数值。已有知识中的10篇背景/方法参考资料在本次未能联网复核出版者记录或全文。\n',
 '## 最重要的判断\n',
 '课题应围绕“保存成效—关联因素—后续衰退识别—分类管护”形成闭环。是否超越已有河西走廊梭梭研究，当前尚不能确定；需先补齐区域直接证据。方法上的潜在改进是独立林分重复、同龄/立地匹配、同株随访、按林分验证，以及把建议做成带条件和监测要求的地块清单。\n',
 '梭梭林现存密度、同批保存比例、个体健康、自然更新和防护功能分别回答不同问题。应避免随意合并成一个总分，也应避免把水分等解释变量写入健康标签后，再用同一变量解释该标签。\n',
 '敦煌属于酒泉市：可单列子区，但不与酒泉总计重复。备选地区与地块需要管理单位记录及现场核实；本稿没有已确认的样地清单。\n',
 '## PPT逐页讲稿\n']
for i,s in enumerate(SLIDES,1):
 md += [f"### {i:02d} {s['title'].replace(chr(10),'：')}\n",s['sub']+'\n']
 for h,b in s['cards']:md.append(f"- **{h}**：{b.replace(chr(10),' ')}")
 if s['table']:
  hdr,rows=s['table'];md.extend(['','| '+' | '.join(hdr)+' |','| '+' | '.join(['---']*len(hdr))+' |'])
  for row in rows:md.append('| '+' | '.join(c.replace('\n','<br>') for c in row)+' |')
 md += ['','讲解与边界：'+s['note'],'','来源/状态：'+s['source'],'']
md += ['## 参考文献与本课题用途\n']
for r in REFS:
 md += [f"### {r['id']} {r['author']} ({r['year']})",r['title']+'. '+r['journal']+'.',f"DOI：https://doi.org/{r['doi']}",'',
        f"研究对象：{r['object']}。",f"问题与方法：{r['question']} {r['method']}。",f"主题归纳：{r['result']}",f"局限：{r['limit']}",f"借鉴：{r['use']}",
        '核验状态：依据已有知识整理；本次未联网复核题录及全文。','']
md += ['## 区域文献补齐后如何更新\n',
 '1. 以数据库正式题录和论文首页核对作者、题名、年份、DOI、区域及物种。',
 '2. 人工林/天然林、梭梭/白梭梭分开；综述与原始研究分开。',
 '3. 提取独立林分数、样方数、个体数和采样时间，不能只记“样本量”。',
 '4. 按现状描述、匹配解释、随访验证及措施试验标记证据层次，同时记录缺失信息。',
 '5. 只有可比指标和独立研究足够时才考虑效应量合并；不同含水率单位、土层和季节不得直接平均。',
 '6. 用直接文献实际覆盖来修订创新点；如已有研究完成闭环，就聚焦更严格验证或一个管护试验。',
 '7. 替换PPT中的待核验页和预期成果页，不添加虚构图表。\n',
 '## 实施前需要确认的条件\n',
 '造林档案与地块名录能否获得；能否在同一季节做年度复测；可用人力、车辆与仪器；能否进入林地并保留植株标签；是否具备一个小规模管理试验的条件。若只能一次调查，应把目标改为现状诊断与脆弱性分型。\n',
 '## 交付与使用\n',
 'PPT为可编辑矢量文字、表格与图形，并含逐页备注；Excel包含10篇背景/方法资料、区域待补主题、12项指标、检索式和现场方案。PDF供快速预览。正式提交开题报告前，应核实参考资料并补齐区域直接文献。']
(OUT/'文献梳理说明与逐页讲稿.md').write_text('\n'.join(md),encoding='utf-8')
(OUT/'slides_manifest.json').write_text(json.dumps(SLIDES,ensure_ascii=False,indent=2),encoding='utf-8')

# Structural checks; visual rendering is performed separately in LibreOffice.
check=Presentation(pptx_path)
assert len(check.slides)==len(SLIDES)
assert all(sl.has_notes_slide for sl in check.slides)
assert not any('TODO' in sh.text for sl in check.slides for sh in sl.shapes if sh.has_text_frame)
assert load_workbook(xlsx_path).sheetnames==['阅读说明','背景与方法文献','区域文献待补','评价指标与口径','检索与核验','现场与分析方案']
print(json.dumps({'slides':len(check.slides),'references':len(REFS),'metric_rows':len(METRICS),'pptx':str(pptx_path),'xlsx':str(xlsx_path)},ensure_ascii=False))
